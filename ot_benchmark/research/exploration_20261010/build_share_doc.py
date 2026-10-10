from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'OT项目双方向研究规划与协作任务书_20261010.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.8)
sec.footer_distance = Inches(0.3)

def font(style, size, bold=False, east='宋体'):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key, val in [('ascii', 'Calibri'), ('hAnsi', 'Calibri'), ('eastAsia', east)]:
        rf.set(qn('w:' + key), val)

font(doc.styles['Normal'], 11)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.16
normal.space_after = Pt(6)
normal.widow_control = True
font(doc.styles['Title'], 22, True, '微软雅黑')
doc.styles['Title'].paragraph_format.space_after = Pt(10)
for name, size in [('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11)]:
    font(doc.styles[name], size, True, '微软雅黑')
    pf = doc.styles[name].paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(5)
    pf.keep_with_next = True
font(doc.styles['Subtitle'], 11, False, '宋体')
font(doc.styles['Caption'], 10)
for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
    ppr = doc.styles[name].element.pPr
    if ppr is not None:
        for el in list(ppr):
            if el.tag == qn('w:pBdr'):
                ppr.remove(el)

doc.core_properties.title = 'OT项目双方向研究规划与协作任务书'
doc.core_properties.subject = 'OT算法研究与AIGC应用研究的候选方案和并行协作'
doc.core_properties.author = '项目组'
doc.core_properties.keywords = '最优传输 OT AIGC 评测 研究规划'

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run._r.addnext(field)

def p(text, lead=None, small=False):
    para = doc.add_paragraph()
    if lead:
        para.add_run(lead + ' ').bold = True
    r = para.add_run(text)
    if small:
        r.font.size = Pt(10)
    return para

def h(text, level=2):
    return doc.add_heading(text, level)

def page(title):
    doc.add_page_break()
    h(title, 1)

def link(para, label, url):
    rid = para.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyp = OxmlElement('w:hyperlink')
    hyp.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '164E79')
    rp.append(color)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '20')
    rp.append(sz)
    r.append(rp)
    t = OxmlElement('w:t')
    t.text = label
    r.append(t)
    hyp.append(r)
    para._p.append(hyp)

def table(headers, rows, widths):
    tb = doc.add_table(rows=1, cols=len(headers))
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb.autofit = False
    for c, width in zip(tb.columns, widths):
        c.width = Inches(width)
    pr = tb._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        for attr, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            el.set(qn('w:' + attr), val)
        borders.append(el)
    pr.append(borders)
    for vals in [headers] + rows:
        row = tb.rows[0] if vals is headers else tb.add_row()
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        is_head = vals is headers
        if is_head:
            repeat = OxmlElement('w:tblHeader')
            row._tr.get_or_add_trPr().append(repeat)
        for j, (cell, val) in enumerate(zip(row.cells, vals)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top', '90'), ('bottom', '90'), ('left', '100'), ('right', '100')]:
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), value)
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tcpr.append(margins)
            if is_head:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E7EDF2')
                tcpr.append(shade)
            cell.text = val
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.line_spacing = 1.08
            for rr in cp.runs:
                rr.font.size = Pt(10)
                rr.bold = is_head
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return tb

doc.add_paragraph('OT项目双方向研究规划与协作任务书', 'Title')
doc.add_paragraph('协作讨论与首轮探索安排    2026年10月10日', 'Subtitle')
p('我们继续以原来的OT评测平台为基础，分两条研究线寻找创新点。建议先推进1A和2A，其他方案先做小规模筛查。以下六条都是待验证的候选；目前还没有证明新算法有效，也没有验证AIGC最终效果提高。')
h('项目目标与分工')
p('大创总体目标仍是回答不同数据条件下该选什么算法，以及怎样公平、自动地比较算法。两条研究线从平台发现的问题出发，为项目补充有深度的研究成果，不等于正式修改立项方向。')
table(['研究线', '研究目标', '建议分工'], [
    ['方向1', '设计新的OT求解或近似方法，研究精度与时间内存的平衡', '合作成员推进'],
    ['方向2', '改造具体应用中的匹配或引导环节，提高最终生成编辑效果', '我们继续推进'],
], [0.7, 4.4, 1.8])
h('六条候选与优先级')
table(['编号', '候选方案', '首轮定位'], [
    ['1A', '面向分簇数据的簇间慢模态校正', '方向1主线'],
    ['1B', '保留必要跨簇通道的自适应稀疏OT', '备选筛查'],
    ['1C', '误差预算驱动的因子化计划与局部细化', '备选筛查'],
    ['2A', '对象完整性约束的空间关系OT引导', '方向2主线'],
    ['2B', '语义区域守恒的参考图风格传输', '备选筛查'],
    ['2C', '编辑区域与保留区域分工的OT轨迹引导', '备选筛查'],
], [0.7, 4.4, 1.8])
p('方向2先用成熟求解器，不等待方向1成功。将来接入新求解器时，分别检验应用建模和求解器的贡献。已有“模块A”是成熟算法对照模块，不是新研究方向1，不能直接覆盖。')
p('阅读建议：方向1成员重点看第2页和第5页；方向2重点看第3至5页。两边选中主线后，先补齐最接近3至5篇论文的差异表，再实现最小原型。', small=True)

page('方向1 新OT算法的三条候选')
p('现有平台已有生成数据、统一运行、质量诊断和结果追溯。在特定1D双簇协议下，成熟基线模块完成432次调用、19项测试；普通Sinkhorn在24/48条件达标，稳定化和正则延续各30/48，按三个重复均通过计。近平衡非零簇间质量差是共同困难，但还不能推广为高维结论。')
h('1A 簇间慢模态校正')
p('簇内匹配已稳定时，少量跨簇质量可能移动很慢。尝试在固定最终正则参数的对偶空间中，构造簇常数粗空间，求小规模粗校正方向，再做线搜索和细层Sinkhorn更新。需要处理势的平移自由度与近奇异系统。', '核心设想')
p('创新候选是自动识别真正慢的跨簇模式，并由原始目标误差决定校正时机。Sinkhorn加Newton已有工作[2][3]，因此必须比较普通粗初始化、固定周期校正、随机粗空间、SNS或准Newton和正则延续。', '区别与对照')
p('历史1D族只用于开发；独立测试改变簇数、质量差、簇内宽度、重叠和初始化，再扩展真实高维特征。在原始成本和相同精度下比较达标时间、可靠性及全部开销。若收益仅来自已知双簇标签，或扰动后消失，降级为机制研究。仍构造完整矩阵时不声称省内存。', '验证与止损')
h('1B 自适应稀疏OT')
p('近邻截断可能删掉很小但必要的跨簇运输。尝试从局部近邻和少量跨簇边开始，用质量缺额、对偶约化成本与遗漏质量界补边，并安排分块全问题审核。', '核心设想')
p('稀疏核、多尺度和块筛选已有[1][9][10]；候选增量在质量可行性与原始误差共同驱动的补边预算。图连通不保证给定边缘质量可行；正熵OT的有限稀疏计划是近似，稀疏子问题的gap不能证明原始问题达标。', '区别与风险')
p('比较固定近邻、固定阈值、成熟筛选与多尺度方法。必须计入遗漏边审核、补边、回退的时间和内存；避免先构造完整成本矩阵。若全局审核抵消收益，或边数趋近稠密，则停止。', '验证与止损')
h('1C 自适应因子化计划与局部细化')
p('以真正的低秩运输计划起步，根据剩余误差决定增加全局秩、拆分因子或细化局部结构，保持非负与边缘约束。候选贡献是误差预算下的容量分配规则；固定低秩、低秩加稀疏及层次细化已有[5][6][7]。', '核心设想')
p('对比多档固定秩、单纯增秩与兼容场景下的层次方法，画误差—时间—内存前沿。低秩问题的驻点不等于原始OT全局解。只有成本算子也避免稠密存储、且不重建完整计划，才谈端到端省内存；若只是多试几档秩，保留为平台扩展。', '验证与止损')

page('方向2 优先尝试空间关系OT引导')
h('2A 对象完整性约束的空间关系OT引导')
p('场景是“红猫在蓝狗左边”“小杯子在大桌子上方”等多对象文生图。改动去噪过程中注意力的空间引导，目标是对象都出现、关系与属性正确，同时保持图像质量。该路线主要是已有OT改为任务适配的OT；另设非OT引导对照。')
p('使用真正的源—目标位置代价，按对象置信度设置匹配容量。可以不强制迁移不确定背景，但要保留对象存在和有效匹配质量下限，防止非平衡OT通过删除难对象降低损失。最初只改一个OT建模机制，固定生成骨干、噪声和更新预算。', '候选机制')
h('先完成STORM的论文代码一致性审计')
p('STORM已经用OT控制空间布局[11]。静态读取其公开pipeline.py发现，一条成本构造路径将同一个行成本重复到所有列。在精确平衡、固定边缘质量下，这类成本不能区分运输匹配路径；而论文补充材料描述了含位置间距离的混合成本。')
p('这是尚待端到端复核的一致性疑点，不代表所有损失项都退化，也不否定论文报告的效果。有限迭代及分母稳定项可能使计划未精确满足边缘约束，不能直接套用精确条件下的结论。')
p('首轮至少比较原版STORM、纯行成本解析对照、论文一致的混合位置成本、稳定化且质量匹配的求解版本，以及候选方法。补齐成本或稳定化属于复现修正；新方案只有超过修正后的强基线，才讨论研究贡献。')
h('实验与最终验收')
p('先用同一模型复现STORM、无引导生成和非OT注意力引导。开发40条提示×4个种子，冻结至少200条独立提示×4个种子做初筛；通过后扩展完整官方协议。上述规模是探索建议，不是已经完成的实验。')
p('最终主指标覆盖全体输出的对象存在与关系正确率，例如VISOR的无条件关系指标；缺对象必须计失败。同时报告属性绑定、复杂组合及T2I-CompBench完整任务分项[18]，保留固定评估器和盲审质量。不能只报内部OT损失或挑选空间子集。', '主指标')
p('分别消融几何成本、对象质量下限、置信度、平衡与非平衡匹配、固定与自适应引导。提供相同梯度步数的非OT质心或布局损失对照；所有方法使用同样的布局信息。', '公平对照')
p('若空间分提高但对象存在、属性或质量明显下降；只胜公开代码而不胜论文一致版本；或者优势仅来自更多更新，均不保留为主成果。', '止损条件')

page('方向2 两条备选应用路线')
h('2B 语义区域守恒的参考图风格传输')
p('任务是只改变衣服、动物或建筑的参考纹理，保留脸、轮廓和背景。可在冻结的Attention Distillation或特征注入环节[12]，将softmax或最近邻匹配改为带语义区域容量及拒配的OT。不同面积、缺失语义区域不强行一一匹配；不可靠匹配保留原内容特征。')
p('候选增量是不等面积、不完全重叠区域下的约束与不确定性处理。UOT替代注意力、颜色OT及区域风格控制已有[13][14][15]，不能直接称为新颖。应保留原Q/K尺度和残差结构，先核对匹配接口，而非盲目替换整层注意力。', '创新边界')
p('比较原非OT匹配、简单mask门控、平衡OT、普通UOT和候选区域OT，使用同一自动区域估计流程。真值mask仅用于评估，oracle实验单列。开发30组内容—参考—区域组合×3个种子，再冻结至少100组独立组合，防止内容或参考图跨集合泄漏。', '最小实验')
p('主指标是在区域外保护门槛下的区域风格达成；同时报告区域外LPIPS/MSE、身份轮廓和盲审。完全不编辑与修改背景均不能算成功。若与简单mask门控无差异、复制参考对象内容、或最近论文已覆盖机制，则停止。RegionRoute含训练与LoRA，跨系统比较必须单列资源差异。', '验收与止损')
h('2C 编辑区域与保留区域分工的OT轨迹引导')
p('任务是红车变蓝、给人加眼镜或替换物体，保留非编辑区。可在FlowEdit或Rectified Flow编辑的速度场修正环节[16][17]，显式建立patch或token匹配：保留区强调几何语义一致，编辑区放松对应或允许部分匹配，不确定区降低引导，限制对原更新的扰动。')
p('OT-RF已做闭式位移和时间调度引导；所读方法并非离散Sinkhorn层，不能把方案描述为“替换它的Sinkhorn”。加入token级OT是新的建模假设，必须证明对应关系与最终效果有意义，仅加mask或改时间权重不足以说明贡献。', '创新边界')
p('比较FlowEdit、OT-RF、同区域信息的mask化OT-RF、简单局部插值、标准平衡或部分OT，以及候选。按PIE-Bench任务类型划分开发测试，先筛查80个独立案例，再扩展完整协议。同骨干、源图、噪声、步数和区域信息配对；评估真值mask不得泄漏到候选。', '最小实验')
p('主指标是背景保护门槛下的编辑达成率，并报告编辑区目标一致性、区域外LPIPS和身份结构。无实际编辑计失败，重建更好不等于编辑成功。若保护背景却压制编辑、与mask基线无差异，或优势来自更强模型而非机制，则停止。', '验收与止损')

page('共同实验协议与协作安排')
h('评测与接口要求')
p('方向1比较相同原始目标、正则参数和质量门槛下的时间、内存与达标率；先试目标间隙1e-6和1e-8两档。方向2预先冻结最终任务主指标、不可牺牲的质量指标和资源预算。OT求得更精确不保证图像更好，内部误差只能解释结果。')
p('所有预处理、聚类、成本构造、求解、审核和回退计入端到端时间。GPU同步计时，CPU与GPU峰值内存分别报告；模型载入、训练和特征缓存另列。未运行、超时、无效、质量失败和成功完整保留，先报达标率，再比较共同达标实例耗时。')
p('开发、验证、测试独立。以提示或图像案例为统计单位，先汇总同案例的随机种子，再按案例做配对bootstrap。全量保存生成图和负例，不以少量图片的FID或精选展示图证明总体提高。')
p('保留原距离标量插件接口，新增可返回计划、因子或映射的研究接口，并记录目标类型、原始成本、正则、势、验收原因、分阶段耗时、内存、种子与配置代码模型哈希。非平衡、部分OT等目标需要自己的验收，不能套平衡EOT诊断。')
h('首轮工作安排')
table(['阶段', '方向1', '方向2'], [
    ['复现审计\n约2至3工作日', '读1A近邻全文；重放失败；统一精度强基线', '审计STORM；跑通最终指标；检查2B与2C代码数据可用性'],
    ['最小原型\n约3至5工作日', '实现1A最小校正；筛查1B与1C的关键风险', '实现2A一个机制；2B与2C各做小样本筛查'],
    ['正式验证\n实测后排期', '保留一条；跨结构规模与真实数据验证', '保留一个应用；冻结独立测试与完整协议'],
], [1.3, 2.8, 2.8])
p('上述是工作窗口，不是效果或论文发表保证。正式排期由每实例耗时、方法数、样本数和显存实测确定。每条路线都需要给出支持结果与失败结果；仅内部损失改善的应用方案不进入正式主线。')
h('算力与协作交付')
p('算力按RTX 5070 Ti规划，必要时使用云平台。先核对可用显存、驱动、CUDA、权重和模型许可；付费云使用前确认预算。第一版冻结预训练模型，不以大规模训练作为起点。')
p('两个方向使用独立目录、环境、开发测试清单、种子命名空间与分支，保留模块A。方向1建议交付目标与算法定义、可复现代码、统一精度对照、原始记录、失败边界和论文差异表。我们负责方向2实现与最终效果评测，并汇总双方结论；合并前共同复核，避免覆盖对方修改。')
p('目前完成的是方案、文献依据与协作规划。新方案查新、AIGC端到端复现、算法实现和最终得分提高都还待执行。')

def ref(number, title, venue, why, url, label='论文或官方页面'):
    para = p(f'[{number}] {title}。{venue}。', small=True)
    para.paragraph_format.keep_with_next = True
    para = p(why + ' ', small=True)
    para.paragraph_format.space_after = Pt(9)
    link(para, label, url)

page('参考论文 OT求解与近似')
p('以下链接来自前期检索与原文或官方页面核验。部分工作只核对了摘要和会议记录，尚未穷尽全文查新；各路线选定后仍需补充最接近论文的算法与假设差异。预印本与已发表论文分别标注。', small=True)
ref(1, 'Schmitzer — Stabilized Sparse Scaling Algorithms for Entropy Regularized Transport Problems', 'SIAM Journal on Scientific Computing 2019；预印本2016', '稳定化、截断、多尺度与正则延续基线；原文及项目原文笔记。', 'https://arxiv.org/abs/1610.06519')
ref(2, 'Tang et al. — Accelerating Sinkhorn algorithm with sparse Newton iterations', 'ICLR 2024', '1A的Newton强近邻；核对会议页面与摘要。', 'https://proceedings.iclr.cc/paper_files/paper/2024/hash/3a819a53408e20b75d1954bf617ccc0a-Abstract-Conference.html')
ref(3, 'Wang and Qiu — The Sparse-Plus-Low-Rank Quasi-Newton Method for Entropic-Regularized Optimal Transport', 'ICML 2025', '稀疏加低秩Hessian近似已有；不是低秩运输计划。核对会议页面与摘要。', 'https://proceedings.mlr.press/v267/wang25ck.html')
ref(4, 'Qiu — cuRegOT A GPU-Accelerated Solver for Entropic-Regularized Optimal Transport', 'ICML 2026', 'GPU高性能基线；核对会议页面与摘要。', 'https://proceedings.mlr.press/v306/qiu26y.html')
ref(5, 'Scetbon Cuturi and Peyré — Low-Rank Sinkhorn Factorization', 'ICML 2021', '1C真正低秩耦合的基础；核对会议页面与摘要。', 'https://proceedings.mlr.press/v139/scetbon21a.html')
ref(6, 'Halmos et al. — Hierarchical Refinement Optimal Transport to Infinity and Beyond', 'ICML 2025', '低秩与层次细化近邻；assignment设定需区别非均匀质量。核对摘要。', 'https://proceedings.mlr.press/v267/halmos25a.html')
ref(7, 'Liu et al. — Approximating Optimal Transport via Low-rank and Sparse Factorization', '2021预印本', '低秩加稀疏已有；核对作者页与摘要。', 'https://arxiv.org/abs/2111.06546')
ref(8, 'Hsieh et al. — SinkSLOT Sinkhorn via Sparse Lifted Optimal Transport', '2026年8月28日预印本', '稀疏先验可能改变目标，不按同一EOT证书混排。核对作者页与摘要。', 'https://arxiv.org/abs/2608.28262')
ref(9, 'Ye et al. — FlashSinkhorn 2 Block-Sparse Entropic Optimal Transport', '2026年10月1日预印本', '1B块筛选最新近邻，尚非同行评审结论。核对作者页与摘要。', 'https://arxiv.org/abs/2610.02395')
ref(10, 'GeomLoss — Multiscale Sinkhorn and kernel truncation', '官方实现文档', '多尺度与截断基线；比较时统一目标、正则和debias选项。', 'https://www.kernel-operations.io/geomloss/_auto_examples/sinkhorn_multiscale/plot_kernel_truncation.html', '官方文档')

page('参考论文 AIGC应用与评测')
ref(11, 'Han et al. — Spatial Transport Optimization by Repositioning Attention Map for Training-Free Text-to-Image Synthesis', 'CVPR 2025', 'STORM是2A强基线；读取方法、附录与官方代码关键路径。', 'https://arxiv.org/html/2503.22168v1')
ref(12, 'Zhou et al. — Attention Distillation A Unified Approach to Visual Characteristics Transfer', 'CVPR 2025', '2B非OT参考特征转移基线；核对会议索引与作者项目代码。', 'https://arxiv.org/abs/2502.20235')
ref(13, 'Li et al. — Arbitrary style transfer with attentional networks via unbalanced optimal transport', 'IET Image Processing 2022', 'UOT注意力已有，不能作为2B独立新颖性；核对出版社全文3.2与3.3。', 'https://ietresearch.onlinelibrary.wiley.com/doi/full/10.1049/ipr2.12403')
ref(14, 'Wang et al. — COTA Color-guided optimal transport attention for arbitrary style transfer', 'Array 31 101197 2026年9月', '颜色OT与门控近邻；核对作者机构记录，尚未全文复核。', 'https://researchoutput.csu.edu.au/en/publications/cota-color-guided-optimal-transport-attention-for-arbitrary-style/', '作者机构记录')
ref(15, 'Chen et al. — RegionRoute Regional Style Transfer with Diffusion Model', 'CVPR 2026', '区域风格强近邻；读取作者原文方法与指标。', 'https://arxiv.org/html/2602.19254v1')
ref(16, 'Kulikov et al. — FlowEdit Inversion-Free Text-Based Editing Using Pre-Trained Flow Models', 'ICCV 2025', '2C非OT编辑强基线；核对会议索引与官方代码。', 'https://openaccess.thecvf.com/content/ICCV2025/html/Kulikov_FlowEdit_Inversion-Free_Text-Based_Editing_Using_Pre-Trained_Flow_Models_ICCV_2025_paper.html')
ref(17, 'Lupaşcu and Stupariu — Optimal Transport for Rectified Flow Image Editing Unifying Inversion-Based and Direct Methods', 'WACV 2026', 'OT-RF已做轨迹引导；读取作者原文3.3至3.5与算法。', 'https://arxiv.org/html/2508.02363v2')
ref(18, 'Huang et al. — T2I-CompBench A Comprehensive Benchmark for Open-world Compositional Text-to-image Generation', 'NeurIPS 2023 Datasets and Benchmarks', '最终图像组合任务评测；核对官方会议页面。', 'https://papers.neurips.cc/paper_files/paper/2023/hash/f8ad010cdd9143dbb0e9308c093aff24-Abstract-Datasets_and_Benchmarks.html')
h('代码与数据入口')
for label, url in [
    ('OT项目仓库', 'https://github.com/iJay0928/ot-benchmark-'),
    ('STORM官方代码', 'https://github.com/MICV-yonsei/STORM'),
    ('Attention Distillation官方代码', 'https://github.com/xugao97/AttentionDistillation'),
    ('RegionRoute官方代码', 'https://github.com/bwchen05/RegionRoute'),
    ('FlowEdit官方代码', 'https://github.com/fallenshock/FlowEdit'),
    ('T2I-CompBench官方代码', 'https://github.com/Karine-Huang/T2I-CompBench'),
    ('PIE-Bench数据入口', 'https://github.com/cure-lab/PnPInversion'),
]:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(2)
    link(para, label, url)

doc.save(OUT)
print(OUT)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables))
